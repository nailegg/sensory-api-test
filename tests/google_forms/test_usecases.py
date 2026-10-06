import pytest

from app.services.google_drive.client import DriveApiError
from app.services.google_forms import usecases
from app.services.google_forms.client import FormsApiError
from app.services.google_forms.usecases import ANYONE_WITH_LINK_PERMISSION_ID, build_setup_requests


class FakeForms:
    def __init__(self, fail_batch: bool = False):
        self.calls: list[tuple] = []
        self.fail_batch = fail_batch

    async def create(self, title, document_title=None, unpublished=None):
        self.calls.append(("create", title, unpublished))
        return {"formId": "f1"}

    async def batch_update(self, form_id, requests, include_form_in_response=False, required_revision_id=None):
        self.calls.append(("batch_update", form_id, requests))
        if self.fail_batch:
            raise FormsApiError(400, {"error": {"message": "Invalid grading", "status": "INVALID_ARGUMENT"}})
        return {"replies": [{} for _ in requests]}

    async def set_publish_settings(self, form_id, is_published, is_accepting_responses):
        self.calls.append(("publish", form_id, is_published, is_accepting_responses))
        return {}

    async def get(self, form_id, fields=None):
        return {"formId": form_id, "info": {"title": "설문"}, "publishSettings": {"publishState": {"isPublished": True, "isAcceptingResponses": True}}}


class FakeDrive:
    def __init__(self, fail_emails: set[str] = frozenset(), delete_status: int | None = None):
        self.calls: list[tuple] = []
        self.fail_emails = fail_emails
        self.delete_status = delete_status

    async def get_file(self, file_id, fields=None):
        return {"id": file_id, "name": "설문", "mimeType": "application/vnd.google-apps.form"}

    async def move_file(self, file_id, to_folder_id, from_folder_id=None):
        self.calls.append(("move", file_id, to_folder_id))
        return {}

    async def trash_file(self, file_id):
        self.calls.append(("trash", file_id))
        return {}

    async def share_as_responder(self, file_id, email=None, notify=False):
        if email in self.fail_emails:
            raise DriveApiError(400, {"error": {"message": "bad", "errors": [{"reason": "invalidSharingRequest"}]}})
        self.calls.append(("share", file_id, email))
        return {"id": f"perm-{email}", "view": "published"}

    async def delete_permission(self, file_id, permission_id):
        if self.delete_status:
            raise DriveApiError(self.delete_status, {"error": {"message": "x", "errors": [{"reason": "notFound"}]}})
        self.calls.append(("delete_permission", file_id, permission_id))


def test_setup_requests_order_quiz_before_items():
    reqs = build_setup_requests(description="설명", items=[{"title": "a"}, {"title": "b"}], email_collection="VERIFIED", quiz=True)
    assert reqs[0] == {
        "updateSettings": {
            "settings": {"quizSettings": {"isQuiz": True}, "emailCollectionType": "VERIFIED"},
            "updateMask": "quizSettings.isQuiz,emailCollectionType",
        }
    }
    assert reqs[1]["updateFormInfo"]["updateMask"] == "description"
    assert [r["createItem"]["location"]["index"] for r in reqs[2:]] == [0, 1]


def test_setup_requests_rejects_unknown_email_collection():
    with pytest.raises(ValueError):
        build_setup_requests(email_collection="ALWAYS")


@pytest.mark.asyncio
async def test_create_form_starts_unpublished_then_publishes_after_items():
    forms, drive = FakeForms(), FakeDrive()
    info = await usecases.create_form(forms, drive, "설문", folder_id="folder1", items=[{"title": "q"}])
    kinds = [c[0] for c in forms.calls]
    assert forms.calls[0] == ("create", "설문", True)  # 기본값은 게시 상태라 항상 미게시로 만든다
    assert kinds == ["create", "batch_update", "publish"]
    assert ("move", "f1", "folder1") in drive.calls
    assert info.published is True


@pytest.mark.asyncio
async def test_create_form_unpublished_for_scheduled_open():
    forms = FakeForms()
    await usecases.create_form(forms, FakeDrive(), "수업 퀴즈", items=[{"title": "q"}], publish=False)
    assert "publish" not in [c[0] for c in forms.calls]


@pytest.mark.asyncio
async def test_create_form_trashes_empty_form_when_items_fail():
    forms, drive = FakeForms(fail_batch=True), FakeDrive()
    with pytest.raises(FormsApiError):
        await usecases.create_form(forms, drive, "설문", folder_id="folder1", items=[{"title": "q"}])
    assert ("trash", "f1") in drive.calls
    assert "publish" not in [c[0] for c in forms.calls]


@pytest.mark.asyncio
async def test_restrict_responders_removes_link_access_after_adding():
    drive = FakeDrive(fail_emails={"bad@example.com"})
    results = await usecases.restrict_responders(drive, "f1", ["a@example.com", "bad@example.com"])
    assert [(r.email, r.ok) for r in results] == [("a@example.com", True), ("bad@example.com", False), (None, True)]
    assert drive.calls[-1] == ("delete_permission", "f1", ANYONE_WITH_LINK_PERMISSION_ID)


@pytest.mark.asyncio
async def test_restrict_responders_keeps_link_access_when_nobody_added():
    drive = FakeDrive(fail_emails={"bad@example.com"})
    results = await usecases.restrict_responders(drive, "f1", ["bad@example.com"])
    assert all(c[0] != "delete_permission" for c in drive.calls)
    assert [r.email for r in results] == ["bad@example.com"]


@pytest.mark.asyncio
async def test_restrict_responders_rerun_is_safe_when_link_already_removed():
    results = await usecases.restrict_responders(FakeDrive(delete_status=404), "f1", ["a@example.com"])
    assert results[-1].email is None and results[-1].ok


# ---------- 3·6·9~13 ----------
import csv  # noqa: E402
import io  # noqa: E402

from app.core.models import FormSubmission  # noqa: E402
from app.services.google_forms.usecases import (  # noqa: E402
    Member,
    PeerGroup,
    member_labels,
    peer_review_items,
    response_table,
    summarize,
    summarize_peer_reviews,
)

GRID_QUESTIONS = [
    {"id": "c1a", "title": "기여도 [가]", "kind": "grid_row", "options": ["1", "2", "3"], "point_value": None, "group_title": "기여도", "row_title": "가"},
    {"id": "c1b", "title": "기여도 [나]", "kind": "grid_row", "options": ["1", "2", "3"], "point_value": None, "group_title": "기여도", "row_title": "나"},
]


def _sub(rid, email, answers, when="2026-10-05T10:00:00Z", score=None):
    return FormSubmission(id=rid, form_id="f1", respondent_email=email, submitted_at=when, answers=answers, total_score=score)


def test_summarize_counts_other_and_average():
    questions = [
        {"id": "ch", "title": "역할", "kind": "choice_checkbox", "options": ["기획", "개발", "(기타)"], "point_value": None},
        {"id": "sc", "title": "만족도", "kind": "scale", "options": ["1", "2", "3"], "point_value": None},
        {"id": "tx", "title": "의견", "kind": "paragraph", "options": [], "point_value": None},
    ]
    subs = [
        _sub("r1", "a@x.com", {"ch": ["기획", "개발"], "sc": ["3"], "tx": ["좋음"]}),
        _sub("r2", "b@x.com", {"ch": ["디자인"], "sc": ["1"]}),
    ]
    by_id = {s.id: s for s in summarize(questions, subs)}
    assert by_id["ch"].counts == {"기획": 1, "개발": 1, "(기타)": 1}
    assert by_id["ch"].other_answers == ["디자인"]
    assert by_id["sc"].counts == {"1": 1, "2": 0, "3": 1} and by_id["sc"].average == 2.0
    assert by_id["tx"].text_answers == ["좋음"] and by_id["tx"].response_count == 1


def test_peer_reviews_exclude_self_with_reviewees():
    reviewees = {"c1a": "ga@x.com", "c1b": "na@x.com"}
    subs = [_sub("r1", "ga@x.com", {"c1a": ["3"], "c1b": ["2"]}), _sub("r2", "NA@x.com", {"c1a": ["1"], "c1b": ["3"]})]
    scores = {(p.reviewee, p.criterion): p for p in summarize_peer_reviews(GRID_QUESTIONS, subs, reviewees)}
    assert scores[("ga@x.com", "기여도")].average == 1.0 and scores[("ga@x.com", "기여도")].self_excluded == 1
    assert scores[("na@x.com", "기여도")].average == 2.0  # 대소문자 무시로 자기 평가 3점 제외
    no_map = {(p.reviewee, p.criterion): p for p in summarize_peer_reviews(GRID_QUESTIONS, subs)}
    assert no_map[("가", "기여도")].average == 2.0  # reviewees 없으면 행 제목 기준, 자기 평가 판별 불가


def test_member_labels_disambiguate_same_name():
    assert member_labels([Member("김", "kim1@x.com"), Member("김", "kim2@x.com"), Member("이", "lee@x.com")]) == ["김 (kim1)", "김 (kim2)", "이"]
    items = peer_review_items([Member("가", "a@x.com")], ["기여도", "협업"], ["1", "2"])
    assert [i["title"] for i in items] == ["기여도", "협업"]
    assert items[0]["questionGroupItem"]["questions"][0] == {"required": True, "rowQuestion": {"title": "가"}}


def test_response_table_quiz_columns_and_timezone():
    questions = [{"id": "sid", "title": "학번", "kind": "text", "options": [], "point_value": None},
                 {"id": "q", "title": "404?", "kind": "choice_checkbox", "options": [], "point_value": 2}]
    rows = response_table(questions, [_sub("r1", "a@x.com", {"sid": ["0123"], "q": ["가", "나"]}, "2026-10-05T17:38:44Z", 2.0)])
    assert rows[0] == ["제출 시각", "이메일", "점수", "학번", "404?"]
    assert rows[1] == ["2026-10-06 02:38:44", "a@x.com", "2", "0123", "가, 나"]


class FakeFormsResponses(FakeForms):
    def __init__(self, form, pages):
        super().__init__()
        self.form, self.pages, self.list_calls = form, pages, []

    async def get(self, form_id, fields=None):
        return self.form

    async def list_responses(self, form_id, filter=None, page_size=None, page_token=None, fields=None):
        self.list_calls.append({"filter": filter, "page_token": page_token, "fields": fields})
        return self.pages[int(page_token or 0)]


RESP_FORM = {"info": {"title": "퀴즈"}, "items": [
    {"title": "학번", "questionItem": {"question": {"questionId": "sid", "textQuestion": {}}}},
    {"title": "404?", "questionItem": {"question": {"questionId": "q", "grading": {"pointValue": 2}, "choiceQuestion": {"type": "RADIO", "options": [{"value": "a"}]}}}},
]}
PAGES = [
    {"responses": [{"responseId": "r1", "respondentEmail": "A@x.com", "lastSubmittedTime": "2026-10-05T10:00:00.5Z",
                    "answers": {"sid": {"textAnswers": {"answers": [{"value": "0123"}]}}, "q": {"grade": {}, "textAnswers": {"answers": [{"value": "b"}]}}}}],
     "nextPageToken": "1"},
    {"responses": [{"responseId": "r2", "respondentEmail": "a@x.com", "lastSubmittedTime": "2026-10-05T11:00:00Z", "totalScore": 2,
                    "answers": {"q": {"grade": {"score": 2, "correct": True}, "textAnswers": {"answers": [{"value": "a"}]}}}}]},
]


@pytest.mark.asyncio
async def test_collect_paginates_and_returns_cursor():
    forms = FakeFormsResponses(RESP_FORM, PAGES)
    result = await usecases.collect_responses(forms, "f1", since="2026-10-05T00:00:00Z")
    assert [s.id for s in result.submissions] == ["r1", "r2"]
    assert result.cursor == "2026-10-05T11:00:00Z"
    assert forms.list_calls[0]["filter"] == "timestamp >= 2026-10-05T00:00:00Z" and forms.list_calls[1]["page_token"] == "1"


@pytest.mark.asyncio
async def test_submission_status_case_insensitive_and_resubmitted():
    status = await usecases.get_submission_status(FakeFormsResponses(RESP_FORM, PAGES), "f1", ["a@x.com", "b@x.com"])
    assert status.submitted == ["a@x.com"] and status.missing == ["b@x.com"]
    assert status.resubmitted == {"a@x.com": 2} and status.not_in_roster == []


@pytest.mark.asyncio
async def test_quiz_scores_fill_wrong_answers_with_zero():
    scores = await usecases.collect_quiz_scores(FakeFormsResponses(RESP_FORM, PAGES), "f1")
    assert [(s.total_score, s.max_score) for s in scores] == [(0.0, 2.0), (2.0, 2.0)]
    assert scores[0].grades["q"].score == 0 and not scores[0].grades["q"].correct


@pytest.mark.asyncio
async def test_export_csv_bom_and_xlsx_text_cells():
    forms = FakeFormsResponses(RESP_FORM, PAGES)
    exported = await usecases.export_responses(forms, "f1", "csv")
    assert exported.content.startswith(b"\xef\xbb\xbf") and exported.filename == "퀴즈 응답.csv"
    rows = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert rows[1][3] == "0123"
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO((await usecases.export_responses(forms, "f1", "xlsx")).content)).active
    assert ws["D2"].value == "0123" and ws["D2"].number_format == "@"
    with pytest.raises(ValueError):
        await usecases.export_responses(forms, "f1", "pdf")


class FakeSheets:
    def __init__(self):
        self.calls = []

    async def create(self, title, sheet_names=None):
        self.calls.append(("create", title, sheet_names))
        return {"spreadsheetId": "s1", "properties": {"title": title}, "sheets": [{"properties": {"sheetId": 0, "title": "응답"}}]}

    async def get(self, spreadsheet_id, fields=None, ranges=None):
        return {"sheets": [{"properties": {"sheetId": 7, "title": "응답(수정됨)", "index": 0}}]}

    async def clear_values(self, spreadsheet_id, a1_range):
        self.calls.append(("clear", a1_range))
        return {}

    async def update_values(self, spreadsheet_id, a1_range, values, value_input_option="RAW"):
        self.calls.append(("update", a1_range, len(values), value_input_option))
        return {}


@pytest.mark.asyncio
async def test_export_to_sheet_first_then_rerun_clears_only_synsory_columns():
    forms, drive, sheets = FakeFormsResponses(RESP_FORM, PAGES), FakeDrive(), FakeSheets()
    first = await usecases.export_responses_to_sheet(forms, drive, sheets, "f1", folder_id="folder1")
    assert first.created and first.spreadsheet_id == "s1" and first.columns == 5 and first.rows_written == 2
    assert ("move", "s1", "folder1") in drive.calls
    assert sheets.calls[-1] == ("update", "'응답'!A1", 3, "RAW")
    sheets.calls.clear()
    again = await usecases.export_responses_to_sheet(forms, drive, sheets, "f1", spreadsheet_id="s1", previous_columns=7)
    assert not again.created
    assert sheets.calls[0] == ("clear", "'응답(수정됨)'!A:G")  # 지난번 7열까지만. 그 오른쪽 교수자 메모는 남는다


@pytest.mark.asyncio
async def test_create_group_forms_maps_rows_to_emails_and_isolates_failures():
    class GroupForms(FakeForms):
        async def get(self, form_id, fields=None):
            return {"formId": form_id, "info": {"title": "동료평가"}, "items": [{"title": "기여도", "questionGroupItem": {
                "grid": {"columns": {"options": [{"value": "1"}]}},
                "questions": [{"questionId": "x1", "rowQuestion": {"title": "가"}}, {"questionId": "x2", "rowQuestion": {"title": "나"}}]}}]}

    drive = FakeDrive(fail_emails={"bad@example.com"})
    results = await usecases.create_group_forms(
        GroupForms(), drive, "folder1", "1차",
        [PeerGroup("A조", [Member("가", "ga@x.com"), Member("나", "bad@example.com")])], ["기여도"],
    )
    r = results[0]
    assert r.error is None and r.reviewees == {"x1": "ga@x.com", "x2": "bad@example.com"}
    assert [(x.email, x.ok) for x in r.responders] == [("ga@x.com", True), ("bad@example.com", False), (None, True)]


@pytest.mark.asyncio
async def test_create_form_from_template_copies_retitles_publishes_and_restricts():
    class CopyDrive(FakeDrive):
        async def copy_file(self, file_id, name, parent_folder_id=None):
            self.calls.append(("copy", file_id, name, parent_folder_id))
            return {"id": "c1"}

    forms, drive = FakeForms(), CopyDrive()
    result = await usecases.create_form_from_template(
        forms, drive, "orig", "1차 퀴즈 사본", folder_id="folder1", title="1차 퀴즈", responder_emails=["a@x.com"]
    )
    assert drive.calls[0] == ("copy", "orig", "1차 퀴즈 사본", "folder1")
    assert [c[0] for c in forms.calls] == ["batch_update", "publish"]  # 원본(orig)에는 호출하지 않는다
    assert forms.calls[0][1] == "c1" and forms.calls[0][2][0]["updateFormInfo"]["info"]["title"] == "1차 퀴즈"
    assert [(r.email, r.ok) for r in result.responders] == [("a@x.com", True), (None, True)]
